"""生成测试样例文件，用于验证各格式转换是否正常。

不依赖第三方库的部分（csv/json/html/txt/md）直接写文本；
Office 与 PDF 需要依赖装好后由 make_samples 动态生成。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

_OUT = Path(__file__).resolve().parent.parent / "samples"


def _setup_stdout() -> None:
    """统一 UTF-8 输出，避免 Windows 控制台 GBK 编码报错。"""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[union-attr]
        except Exception:
            pass


def _text_files() -> dict[str, str]:
    return {
        "plain.txt": "这是纯文本测试。\n第二行内容。\n",
        "note.md": "# 已是 Markdown\n\n无需转换，用于验证幂等。\n",
        "data.csv": "姓名,年龄,城市\n张三,28,北京\n李四,35,上海\n王五,42,广州\n",
        "config.json": json.dumps(
            {"name": "测试配置", "version": 1, "items": ["a", "b", "c"]},
            ensure_ascii=False,
            indent=2,
        ),
        "page.html": (
            "<!DOCTYPE html><html><head><meta charset='utf-8'>"
            "<title>示例页面</title></head><body>"
            "<h1>一级标题</h1><p>这是一个<strong>加粗</strong>的段落。</p>"
            "<h2>列表</h2><ul><li>项目一</li><li>项目二</li></ul>"
            "<table><tr><th>列A</th><th>列B</th></tr>"
            "<tr><td>1</td><td>2</td></tr></table>"
            "</body></html>"
        ),
        "feed.xml": (
            "<?xml version='1.0' encoding='utf-8'?>"
            "<root><item><title>标题</title><value>内容</value></item></root>"
        ),
    }


def _make_docx_minimal(out: Path) -> None:
    """不依赖 python-docx，直接用 zip + OOXML 构造一个最小可用的 .docx。

    .docx 本质是 zip 容器，含 word/document.xml 即可被解析器读取。
    """
    import zipfile
    from xml.sax.saxutils import escape

    paragraphs = [
        ("Heading1", "Word 测试文档"),
        ("Normal", "这是一个用于验证转换的段落。"),
        ("Normal", "第二段：包含数字 123 和英文 ABC。"),
    ]

    body = []
    for style, text in paragraphs:
        body.append(
            f'<w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr>'
            f"<w:r><w:t>{escape(text)}</w:t></w:r></w:p>"
        )
    # 一个 2x2 表格
    table_rows = [("表头1", "表头2"), ("值1", "值2")]
    rows_xml = ""
    for row in table_rows:
        cells = "".join(
            f"<w:tc><w:tcPr/><w:p><w:r><w:t>{escape(c)}</w:t></w:r></w:p></w:tc>"
            for c in row
        )
        rows_xml += f"<w:tr>{cells}</w:tr>"
    body.append(f"<w:tbl>{rows_xml}</w:tbl>")

    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{''.join(body)}</w:body></w:document>"
    )

    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )

    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/>'
        "</Relationships>"
    )

    with zipfile.ZipFile(out / "demo.docx", "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/document.xml", document_xml)


def _office_files(out: Path) -> list[str]:
    """生成 Office 样例，缺依赖则跳过并报告。"""
    made: list[str] = []

    # DOCX：优先 python-docx，没有就用内置构造器
    try:
        try:
            from docx import Document
        except ImportError:
            _make_docx_minimal(out)
            made.append("demo.docx")
            print("  ✓ demo.docx（内置构造）")
        else:
            doc = Document()
            doc.add_heading("Word 测试文档", level=1)
            doc.add_paragraph("这是一个用于验证转换的段落。")
            doc.add_paragraph("第二段：包含数字 123 和英文 ABC。")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "表头1"
            table.cell(0, 1).text = "表头2"
            table.cell(1, 0).text = "值1"
            table.cell(1, 1).text = "值2"
            doc.save(out / "demo.docx")
            made.append("demo.docx")
            print("  ✓ demo.docx")
    except Exception as exc:
        print(f"  跳过 docx：{exc.__class__.__name__}: {exc}")

    # XLSX
    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "数据表"
        ws.append(["姓名", "年龄", "城市"])
        ws.append(["张三", 28, "北京"])
        ws.append(["李四", 35, "上海"])
        ws2 = wb.create_sheet("第二页")
        ws2.append(["备注"])
        ws2.append(["这是第二个工作表"])
        wb.save(out / "demo.xlsx")
        made.append("demo.xlsx")
    except Exception as exc:
        print(f"  跳过 xlsx：{exc.__class__.__name__}: {exc}")

    # PPTX
    try:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "PowerPoint 测试"
        slide.placeholders[1].text = "副标题内容"
        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        slide2.shapes.title.text = "第二页"
        slide2.placeholders[1].text_frame.text = "要点一"
        for _ in range(1):
            p = slide2.placeholders[1].text_frame.add_paragraph()
            p.text = "要点二"
        prs.save(out / "demo.pptx")
        made.append("demo.pptx")
    except Exception as exc:
        print(f"  跳过 pptx：{exc.__class__.__name__}: {exc}")

    # PDF
    try:
        try:
            from reportlab.pdfgen import canvas
            from reportlab.lib.pagesizes import A4

            c = canvas.Canvas(str(out / "demo.pdf"), pagesize=A4)
            c.setFont("Helvetica", 16)
            c.drawString(72, 780, "PDF Test Document")
            c.setFont("Helvetica", 11)
            c.drawString(72, 750, "This is a sample PDF for conversion testing.")
            c.drawString(72, 730, "Second line with numbers: 12345")
            c.showPage()
            c.setFont("Helvetica", 16)
            c.drawString(72, 780, "Page Two")
            c.save()
            made.append("demo.pdf")
        except ImportError:
            print("  跳过 pdf：未安装 reportlab（不影响主流程，可用真实 PDF 测试）")
    except Exception as exc:
        print(f"  跳过 pdf：{exc.__class__.__name__}: {exc}")

    return made


def main() -> int:
    _setup_stdout()
    _OUT.mkdir(parents=True, exist_ok=True)
    print(f"样例目录：{_OUT}")

    print("\n[1] 文本类文件")
    for name, content in _text_files().items():
        p = _OUT / name
        p.write_text(content, encoding="utf-8")
        print(f"  ✓ {name}")

    print("\n[2] Office / PDF 文件")
    made = _office_files(_OUT)
    if not made:
        print("  （全部跳过，请先安装依赖）")

    print(f"\n完成：共 {len(list(_OUT.iterdir()))} 个样例")
    return 0


if __name__ == "__main__":
    sys.exit(main())
